#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build the concise operator table workbook from dkao_operators_concise.csv."""
from __future__ import annotations

import csv
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT_DIR = os.path.dirname(os.path.abspath(__file__))
CSV = os.path.join(OUT_DIR, "dkao_operators_concise.csv")

HEAD = ["模型", "TP", "算子", "M", "N", "K", "基线µs", "性能µs",
        "加速比", "提升%", "带宽GB/s", "TOPS", "验收状态"]
NUMFMT = {"基线µs": "0.00", "性能µs": "0.00", "加速比": "0.00",
          "提升%": "0", "带宽GB/s": "0.0", "TOPS": "0.0"}

thin = Side(style="thin", color="D0D0D0")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
head_fill = PatternFill("solid", fgColor="1F4E78")
head_font = Font(bold=True, color="FFFFFF")
alt_fill = PatternFill("solid", fgColor="EAF1F8")
warn_fill = PatternFill("solid", fgColor="FFF2CC")


def read_rows():
    rows = list(csv.DictReader(open(CSV, encoding="utf-8-sig")))
    for r in rows:
        for k in ("TP", "M", "N", "K"):
            r[k] = int(r[k])
        for k in ("基线µs", "性能µs", "加速比", "提升%", "带宽GB/s", "TOPS"):
            r[k] = float(r[k])
    return rows


def main():
    rows = read_rows()
    wb = Workbook()

    # ---- sheet 1: concise operator table ----
    ws = wb.active
    ws.title = "算子性能"
    ws.append(HEAD)
    for c in range(1, len(HEAD) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font = head_fill, head_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    for i, r in enumerate(rows, start=2):
        for j, key in enumerate(HEAD, start=1):
            cell = ws.cell(row=i, column=j, value=r[key])
            cell.border = border
            if key in NUMFMT:
                cell.number_format = NUMFMT[key]
                cell.alignment = Alignment(horizontal="right")
            elif key in ("TP", "M", "N", "K"):
                cell.alignment = Alignment(horizontal="right")
            else:
                cell.alignment = Alignment(horizontal="left")
            if r["验收状态"] != "已最终验收":
                cell.fill = warn_fill
            elif i % 2 == 0:
                cell.fill = alt_fill
    widths = {"模型": 12, "TP": 5, "算子": 24, "M": 7, "N": 7, "K": 7,
              "基线µs": 11, "性能µs": 11, "加速比": 9, "提升%": 9,
              "带宽GB/s": 11, "TOPS": 8, "验收状态": 13}
    for j, key in enumerate(HEAD, start=1):
        ws.column_dimensions[get_column_letter(j)].width = widths[key]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEAD))}{ws.max_row}"
    ws.row_dimensions[1].height = 22

    # ---- sheet 2: by model x TP ----
    ws2 = wb.create_sheet("按模型×TP汇总")
    head2 = ["模型", "TP", "算子数", "平均加速比", "最高加速比", "平均带宽GB/s",
             "平均TOPS", "平均性能µs", "其中仅worker验收"]
    ws2.append(head2)
    for c in range(1, len(head2) + 1):
        cell = ws2.cell(row=1, column=c)
        cell.fill, cell.font = head_fill, head_font
        cell.alignment = Alignment(horizontal="center")
        cell.border = border
    groups = {}
    for r in rows:
        groups.setdefault((r["模型"], r["TP"]), []).append(r)
    for i, ((model, tp), rr) in enumerate(sorted(groups.items()), start=2):
        sp = [x["加速比"] for x in rr]
        vals = [model, tp, len(rr), sum(sp) / len(sp), max(sp),
                sum(x["带宽GB/s"] for x in rr) / len(rr),
                sum(x["TOPS"] for x in rr) / len(rr),
                sum(x["性能µs"] for x in rr) / len(rr),
                sum(1 for x in rr if x["验收状态"] != "已最终验收")]
        ws2.append(vals)
        for j in range(1, len(head2) + 1):
            cell = ws2.cell(row=i, column=j)
            cell.border = border
            if j in (4, 5, 6, 7, 8):
                cell.number_format = "0.00"
                cell.alignment = Alignment(horizontal="right")
            elif j in (2, 3, 9):
                cell.alignment = Alignment(horizontal="right")
        if i % 2 == 0:
            for j in range(1, len(head2) + 1):
                ws2.cell(row=i, column=j).fill = alt_fill
    for j, w in enumerate([12, 5, 8, 11, 11, 13, 10, 12, 18], start=1):
        ws2.column_dimensions[get_column_letter(j)].width = w
    ws2.freeze_panes = "A2"

    # ---- sheet 3: notes ----
    ws3 = wb.create_sheet("说明")
    notes = [
        ["指标", "含义"],
        ["硬件/数据类型", "worker29（j20r1n08 / 10.8.160.17 容器 zth_meta）4×K500SM_AI gfx928；INT8 W8A8 GEMM（HIP C++ + DUMMA），输出 bf16"],
        ["基线µs", "固定 Triton(CUDA-graph replay) 基线中位延迟；与优化 kernel 同协议测量"],
        ["性能µs", "最优 variant 的中位延迟（每个 (模型×TP×算子×M) 取记录中的最优已验收结果）"],
        ["加速比 / 提升%", "加速比 = 基线µs ÷ 性能µs；提升% = (基线µs ÷ 性能µs − 1) × 100"],
        ["带宽GB/s", "算法最小字节数 ÷ 性能µs = (M·K + K·N + 4(M+N) + 2·M·N) ÷ t；非实测 HBM 流量"],
        ["TOPS", "逻辑 INT8 运算率 = 2·M·N·K ÷ t"],
        ["验收状态", "『已最终验收』= 该任务串行最终验收全部 shape 通过（final_report.json status=success）；『仅worker验收』= 该任务未产出成功 final_report，取 worker 并行探索阶段已验收最优 median（黄色行）"],
        ["覆盖范围", f"共 {len(rows)} 个算子实例：Hy3（TP4 M16/M4096、TP8 M16/M4096）、MiniMax-M3（TP4 M4096、TP8 M16/M4096）、GLM5.2（TP4 M4096、TP8 M16/M4096）"],
        ["未覆盖", "DeepSeek-V4 默认 42-shape 的早期任务报告在本地已被清空（0 字节），无可用数据；Hy3 TP8 M16+M4096 合并任务、GLM 旧的任务快照因重复或被更新版本覆盖未纳入"],
        ["数据来源", "从各任务工作区记录直接提取（final_report.json / plan.json / shared_baseline/results.json / workers/*/result.json），未做任何新的性能测试"],
    ]
    for i, pair in enumerate(notes, start=1):
        ws3.append(pair)
        for j in (1, 2):
            cell = ws3.cell(row=i, column=j)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if i == 1:
                cell.fill, cell.font = head_fill, head_font
    ws3.column_dimensions["A"].width = 18
    ws3.column_dimensions["B"].width = 110

    path = os.path.join(OUT_DIR, "dkao_operators_concise.xlsx")
    wb.save(path)
    print("wrote", path)
    print("sheets:", wb.sheetnames, "| data rows:", len(rows))


if __name__ == "__main__":
    main()
